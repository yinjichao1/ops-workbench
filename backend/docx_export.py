"""
docx 回填引擎
读取 templates/简历模板.docx 的 24 张表，把学员填写的字段回填进去，
生成一份格式完全一致的正式 Word 简历。

模板结构索引（已解析确认）：
  Table index  1  -> 1.1 基本信息（r0c0 内嵌套表 8行x4列：姓名/性别/政治面貌...逐格；
                      r0c1：证件照/生活照；表外段落：主修课程、研究方向）
  Table index  2  -> 1.2 教育背景 4x7（最高学历/本科大专/高中）
  Table index  0  -> 提示框（不动）
  Table index  3  -> 1.3 实习经历 #1
  Table index  4  -> 1.3 实习经历 #2
  Table index  5  -> 1.4 项目经历 #1
  Table index  6  -> 1.4 项目经历 #2
  Table index  7  -> 1.5 学生干部经历
  Table index  8  -> 1.6 社会实践经历
  Table index  9  -> 1.7 外语水平
  Table index 10  -> 1.8 计算机水平
  Table index 11  -> 1.9 资格证书
  Table index 12  -> 1.10 获奖情况
  Table index 13  -> 1.11 学术成果
  Table index 14  -> 1.11b 教学经历
  Table index 15  -> 1.12 爱好特长
  Table index 16  -> 1.13 家庭背景
  Table index 17  -> 1.14 求职意向
  Table index 18  -> 1.15 自我评价
  Table index 19+ -> 第二部分参考资料（不动）
"""
import shutil
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import TEMPLATE_PATH, EXPORT_DIR


_FILL_MARK = "data-sig-export-filled"


def _set_cell_text(cell, text, *, font_size=9, bold=False, color=None, align=None):
    """写入单元格：清空原内容后写入，保留单元格格式。

    合并去重必须用 XML 属性标记（tc.set），不能用 Python 对象属性：
    lxml 每次访问同一 XML 节点会生成新代理对象，Python 属性会丢（重复写），
    而已回收代理的内存地址可能被其他节点复用（该写不写 → 内容错位丢失）。
    """
    text = "" if text is None else str(text)
    tc = cell._tc
    if tc.get(_FILL_MARK) == "1":
        return cell.paragraphs[0] if cell.paragraphs else None
    tc.set(_FILL_MARK, "1")

    # 清空
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    p = cell.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def _clear_fill_marks(doc):
    """保存前清除全部写入标记，避免残留在导出文件里。"""
    for el in doc.element.body.iter():
        if el.tag.endswith("}tc") and el.get(_FILL_MARK):
            del el.attrib[_FILL_MARK]


def _fill_merged(table, row_idx, col_idx, text):
    """向可能合并的单元格写入；合并区域内只须写一次。"""
    row = table.rows[row_idx - 1]
    if col_idx - 1 >= len(row.cells):
        return
    _set_cell_text(row.cells[col_idx - 1], text)


def _fill_row(table, row_idx, values, start_col=1):
    """向某行依次填充 values（从 start_col 开始，1-based）"""
    for i, v in enumerate(values):
        _fill_merged(table, row_idx, start_col + i, v)


def _fmt_month(v) -> str:
    """'2003-06' -> '2003年6月'；非该格式原样返回"""
    s = (v or "").strip()
    parts = s.split("-")
    if len(parts) == 2 and parts[0].isdigit() and parts[1].isdigit():
        return f"{int(parts[0])}年{int(parts[1])}月"
    return s


# schema 值 -> 模板 checkbox 选项文字（嵌套表 r2c1）
_POLITICAL_MAP = {
    "群众": "群众", "共青团员": "团员", "中共党员": "党员", "中共预备党员": "预备党员",
}


def _fill_basic_nested(inner, basic: dict):
    """1.1 基本信息嵌套表（Table[1].r0c0 内，8行x4列）逐格回填。

    网格（探针确认）：
      r0: 姓名|[值]      | 性别|[值]        r1: 出生年月|[值] | 民族|[值]
      r2: 政治面貌|□群众 □团员 □党员 □预备党员 | 籍贯|[值]
      r3: 手机号|[值]    | 邮箱|[值]        r4: 身份证号|[值] | 户口性质|□农业 □非农业
      r5: 现居住地|[值]  | 是否留学生|□否 □是
      r6: 通信地址|[值]  | 入学前户口|[值]（表单无此字段，留空）
      r7: 婚姻状况|□未婚 □已婚 | 入党时间|____年__月（党员填）

    值格写入已有空段落（8.5pt，与模板标签一致）；
    checkbox 格为单 run 文本，直接把选中项前的 □ 替换为 ☑（保留原格式）；
    空值一律不动，保留模板原样。
    """
    b = basic or {}

    def cell(r, c):
        return inner.rows[r].cells[c]

    def put(r, c, value):
        v = (value or "").strip()
        if not v:
            return
        c_ = cell(r, c)
        p = c_.paragraphs[0] if c_.paragraphs else c_.add_paragraph()
        run = p.add_run(v)
        run.font.size = Pt(8.5)

    def check(r, c, selected, option_map):
        sel = option_map.get((selected or "").strip())
        if not sel:
            return
        for p in cell(r, c).paragraphs:
            for run in p.runs:
                target = f"□{sel}"
                if target in run.text:
                    run.text = run.text.replace(target, f"☑{sel}")

    put(0, 1, b.get("name"))
    put(0, 3, b.get("gender"))
    put(1, 1, _fmt_month(b.get("birth")))
    put(1, 3, b.get("ethnic"))
    check(2, 1, b.get("political"), _POLITICAL_MAP)
    put(2, 3, b.get("native_place"))
    put(3, 1, b.get("phone"))
    put(3, 3, b.get("email"))
    put(4, 1, b.get("id_card"))
    check(4, 3, b.get("hukou"), {"农业": "农业", "非农业": "非农业"})
    put(5, 1, b.get("residence"))
    check(5, 3, b.get("overseas"), {"否": "否", "是": "是"})
    put(6, 1, b.get("address"))
    check(7, 1, b.get("marriage"), {"未婚": "未婚", "已婚": "已婚"})
    # r7c3 入党时间：党员填写，替换「____年__月（党员填）」占位 run
    jt = _fmt_month(b.get("join_party"))
    if jt:
        for p in cell(7, 3).paragraphs:
            for run in p.runs:
                if "年" in run.text and "月" in run.text:
                    run.text = jt


def _append_after_label(doc, label_prefix: str, value):
    """把值追加到「标签：」正文段落（表外）行尾，保留原段落格式。"""
    v = (value or "").strip()
    if not v:
        return
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t.startswith(label_prefix):
            p.add_run(v)
            break


def _format_edu_cell(row: dict, stage_label: str) -> list:
    """教育背景一行的 7 列"""
    school = (row.get("school") or "").strip()
    major = (row.get("major") or "").strip()
    period = (row.get("period") or "").strip()
    form = "☑全日制" if row.get("fulltime") else "□全日制"
    sno = (row.get("student_no") or "").strip()
    gpa = (row.get("gpa") or "").strip()
    rank = (row.get("rank") or "").strip()

    if stage_label == "高中":
        sno = "（高中不填）"
        gpa_cell = ""
    else:
        parts = []
        if gpa:
            parts.append(f"GPA:{gpa}")
        if rank:
            parts.append(f"排名:{rank}")
        gpa_cell = "  ".join(parts)
    return [stage_label, school, major, period, form, sno, gpa_cell]


def _fmt_period(row: dict) -> str:
    return (row.get("period") or "").strip()


def export_resume(data: dict, photos: dict, out_name: str) -> Path:
    """把一份简历数据回填进模板，输出 docx 路径"""
    doc = Document(str(TEMPLATE_PATH))
    tables = doc.tables

    def T(i):
        return tables[i] if i < len(tables) else None

    # ---- 1.1 基本信息 ----
    t = T(1)
    if t:
        # 左格：嵌套表 8x4 逐格回填 + 表外段落（主修课程/研究方向）
        inner = t.rows[0].cells[0].tables
        if inner:
            _fill_basic_nested(inner[0], data.get("basic"))
        _b = data.get("basic") or {}
        _append_after_label(doc, "主修课程", _b.get("courses"))
        _append_after_label(doc, "研究方向", _b.get("research"))
        # 右列：证件照 + 生活照
        photo_cell = t.rows[0].cells[1]
        for p in list(photo_cell.paragraphs):
            p._element.getparent().remove(p._element)
        p = photo_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("【证件照】")
        r.font.bold = True
        r.font.size = Pt(9)
        idp = (photos or {}).get("id_photo")
        if idp and Path(idp).exists():
            try:
                r2 = p.add_run()
                r2.add_picture(str(idp), width=Cm(2.5))
            except Exception:
                p.add_run("（图片读取失败）").font.size = Pt(8)
        else:
            p.add_run("（未上传）").font.size = Pt(8)
        p2 = photo_cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p2.add_run("【生活照】")
        r3.font.bold = True
        r3.font.size = Pt(9)
        life = (photos or {}).get("life_photo")
        if life and Path(life).exists():
            try:
                r4 = p2.add_run()
                r4.add_picture(str(life), width=Cm(2.5))
            except Exception:
                p2.add_run("（图片读取失败）").font.size = Pt(8)
        else:
            p2.add_run("（未上传）").font.size = Pt(8)

    # ---- 1.2 教育背景 ----
    edu = data.get("education") or {}
    rows = edu.get("rows") or [{}, {}, {}]
    t = T(2)
    if t:
        labels = ["最高学历", "本科/大专", "高中"]
        for i, lab in enumerate(labels):
            src = rows[i] if i < len(rows) else {}
            _fill_row(t, i + 2, _format_edu_cell(src, lab))

    # ---- 1.3 实习经历（模板 2 个块，第 4/5 张表）----
    # 第 3 行「主要工作」与第 4 行「工作成果」为横向合并单元格（col2-col4）
    interns = data.get("internships") or []
    for idx, ti in enumerate((3, 4)):
        t = T(ti)
        if not t:
            continue
        item = interns[idx] if idx < len(interns) else {}
        _fill_merged(t, 1, 2, item.get("company", ""))
        _fill_merged(t, 1, 4, item.get("position", ""))
        _fill_merged(t, 2, 2, _fmt_period(item))
        _fill_merged(t, 2, 4, item.get("location", ""))
        _fill_merged(t, 3, 2, item.get("duties", ""))
        _fill_merged(t, 4, 2, item.get("achievements", ""))

    # ---- 1.4 项目经历（第 6/7 张表）----
    projects = data.get("projects") or []
    for idx, ti in enumerate((5, 6)):
        t = T(ti)
        if not t:
            continue
        item = projects[idx] if idx < len(projects) else {}
        _fill_merged(t, 1, 2, item.get("name", ""))
        _fill_merged(t, 1, 4, item.get("role", ""))
        _fill_merged(t, 2, 2, _fmt_period(item))
        _fill_merged(t, 2, 4, item.get("scale", ""))
        _fill_merged(t, 3, 2, item.get("desc", ""))
        _fill_merged(t, 4, 2, item.get("result", ""))

    # ---- 通用多行表填充器 ----
    def fill_multi(table_idx, items, skip_first_row=True):
        t = T(table_idx)
        if not t:
            return
        data_rows = len(t.rows) - (1 if skip_first_row else 0)
        for i in range(data_rows):
            item = items[i] if i < len(items) else {}
            row = t.rows[(i + 1) if skip_first_row else i]
            yield row, item

    # 1.5 学生干部经历（表 8）：职务/单位社团/起止时间/主要工作
    for row, item in fill_multi(7, data.get("student_leaders") or [], skip_first_row=False):
        _set_cell_text(row.cells[0], item.get("title", ""))
        _set_cell_text(row.cells[1], item.get("org", ""))
        _set_cell_text(row.cells[2], _fmt_period(item))
        _set_cell_text(row.cells[3], item.get("duty", ""))

    # 1.6 社会实践经历（表 9）
    for row, item in fill_multi(8, data.get("practices") or [], skip_first_row=False):
        _set_cell_text(row.cells[0], item.get("name", ""))
        _set_cell_text(row.cells[1], item.get("org", ""))
        _set_cell_text(row.cells[2], _fmt_period(item))
        _set_cell_text(row.cells[3], item.get("content", ""))

    # 1.7 外语水平（表 10）：数据从第 2 行起
    for row, item in fill_multi(9, data.get("languages") or []):
        _set_cell_text(row.cells[0], item.get("exam", ""))
        _set_cell_text(row.cells[1], item.get("score", ""))
        _set_cell_text(row.cells[2], item.get("issuer", ""))
        _set_cell_text(row.cells[3], item.get("date", ""))
        _set_cell_text(row.cells[4], item.get("remark", ""))

    # 1.8 计算机水平（表 11）
    for row, item in fill_multi(10, data.get("computers") or []):
        _set_cell_text(row.cells[0], item.get("exam", ""))
        _set_cell_text(row.cells[1], item.get("score", ""))
        _set_cell_text(row.cells[2], item.get("issuer", ""))
        _set_cell_text(row.cells[3], item.get("date", ""))
        _set_cell_text(row.cells[4], item.get("remark", ""))

    # 1.9 资格证书（表 12）：证书名称/等级/颁发机构/证书编号/获得时间/备注
    for row, item in fill_multi(11, data.get("certificates") or []):
        _set_cell_text(row.cells[0], item.get("name", ""))
        _set_cell_text(row.cells[1], item.get("level", ""))
        _set_cell_text(row.cells[2], item.get("issuer", ""))
        _set_cell_text(row.cells[3], item.get("no", ""))
        _set_cell_text(row.cells[4], item.get("date", ""))
        _set_cell_text(row.cells[5], item.get("remark", ""))

    # 1.10 获奖情况（表 13）
    for row, item in fill_multi(12, data.get("awards") or []):
        _set_cell_text(row.cells[0], item.get("name", ""))
        _set_cell_text(row.cells[1], item.get("issuer", ""))
        _set_cell_text(row.cells[2], item.get("level", ""))
        _set_cell_text(row.cells[3], item.get("date", ""))
        _set_cell_text(row.cells[4], item.get("remark", ""))

    # 1.11 学术成果（表 14）：第1行「类型」横向合并占 col1-2，
    # 数据行未合并：col1=类型 col2=名称 col3=刊物 col4=时间 col5=排名
    for row, item in fill_multi(13, data.get("academics") or []):
        _set_cell_text(row.cells[0], item.get("type", ""))
        _set_cell_text(row.cells[1], item.get("name", ""))
        _set_cell_text(row.cells[2], item.get("journal", ""))
        _set_cell_text(row.cells[3], item.get("date", ""))
        _set_cell_text(row.cells[4], item.get("rank", ""))

    # 1.11b 教学经历（表 15）
    for row, item in fill_multi(14, data.get("teachings") or []):
        _set_cell_text(row.cells[0], item.get("course", ""))
        _set_cell_text(row.cells[1], item.get("audience", ""))
        _set_cell_text(row.cells[2], _fmt_period(item))
        _set_cell_text(row.cells[3], item.get("hours", ""))

    # 1.12 爱好特长（表 16，单格）
    t = T(15)
    if t:
        _set_cell_text(t.rows[0].cells[0], data.get("hobbies") or "无")

    # 1.13 家庭背景（表 17）：关系/姓名/工作单位/职务/政治面貌，父亲母亲配偶三行
    fam = (data.get("family") or {}).get("rows") or [{}, {}, {}]
    t = T(16)
    if t:
        for i, rel in enumerate(["父亲", "母亲", "配偶"]):
            src = fam[i] if i < len(fam) else {}
            row = t.rows[i + 1]
            _set_cell_text(row.cells[0], rel)
            _set_cell_text(row.cells[1], src.get("name", ""))
            _set_cell_text(row.cells[2], src.get("company", ""))
            _set_cell_text(row.cells[3], src.get("position", ""))
            _set_cell_text(row.cells[4], src.get("political", ""))

    # 亲属任职：模板中为独立段落「亲属在电力系统/本企业任职情况」
    _patch_relative_paragraph(doc, data.get("family") or {})

    # 1.14 求职意向（表 18，8 行 2 列）
    intent = data.get("intention") or {}
    t = T(17)
    if t:
        mapping = [
            ("应聘单位/岗位", intent.get("target", "")),
            ("期望工作地点", intent.get("location", "")),
            ("首选考试城市", intent.get("exam_city", "")),
            ("是否服从调剂", intent.get("adjust", "")),
            ("是否接受基层/偏远地区", intent.get("grassroots", "")),
            ("是否考研/考博", intent.get("postgrad", "")),
            ("近期是否出国留学", intent.get("abroad", "")),
            ("其他情况说明", intent.get("other", "") or "无"),
        ]
        for i, (_, val) in enumerate(mapping):
            _set_cell_text(t.rows[i].cells[1], val)

    # 1.15 自我评价（表 19，单格）
    t = T(18)
    if t:
        _set_cell_text(t.rows[0].cells[0], data.get("self_eval") or "")

    # 保存（先清除写入标记）
    _clear_fill_marks(doc)
    safe = _safe_name(out_name)
    ts = _ts()
    out_path = EXPORT_DIR / f"{safe}_{ts}.docx"
    doc.save(str(out_path))
    return out_path


def _patch_relative_paragraph(doc, family: dict):
    """改写「亲属在电力系统/本企业任职情况：」后的选项行"""
    has = family.get("has_relative")
    detail = (family.get("relative_detail") or "").strip()
    # 注意：doc.paragraphs 每次访问都返回新代理列表，不能用 .index(p) 二次查找
    # （身份比较必失败 → ValueError 被 except 吞掉 → 选项行永远改不上）
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        txt = p.text or ""
        if "亲属在电力系统" in txt and "任职情况" in txt:
            nxt = paras[i + 1] if i + 1 < len(paras) else None
            if nxt and ("□无" in nxt.text or "□有" in nxt.text):
                for r in list(nxt.runs):
                    r._element.getparent().remove(r._element)
                if has == "有" and detail:
                    nxt.add_run(f"☑有（{detail}）　□无")
                else:
                    nxt.add_run("☑无　□有（关系/姓名/单位）")
            break


def _safe_name(name: str) -> str:
    bad = '<>:"/\\|?*'
    s = "".join(ch for ch in (name or "简历") if ch not in bad).strip()
    return s or "简历"


def _ts() -> str:
    import time
    return time.strftime("%Y%m%d_%H%M%S")
